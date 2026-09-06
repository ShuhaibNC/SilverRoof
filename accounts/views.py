from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout, get_user_model
from django.contrib.auth.decorators import login_required, user_passes_test
from django.views.decorators.cache import never_cache
from django.db.models import Sum
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from .forms import SalesMemberForm
from orders.models import Order
from django.http import HttpResponse
import openpyxl
from openpyxl.styles import Font
from django.utils import timezone
from datetime import timedelta
import random

User = get_user_model()


# ---------------------------------
# HOME
# ---------------------------------
def home(request):
    return render(request, "home.html")


# ---------------------------------
# CHECK ADMIN
# ---------------------------------
def is_admin(user):
    return user.is_superuser


# ---------------------------------
# ADMIN LOGIN
# ---------------------------------
@never_cache
def admin_login(request):

    if request.method == "POST":

        username = request.POST.get("username")
        password = request.POST.get("password")

        user = authenticate(
            request,
            username=username,
            password=password
        )

        if user is not None and user.is_superuser:

            login(request, user)

            return redirect("admin_dashboard")

        return render(
            request,
            "accounts/admin_login.html",
            {
                "error": "Invalid Admin Username or Password"
            }
        )

    return render(
        request,
        "accounts/admin_login.html"
    )
User = get_user_model()
@never_cache
def admin_forgot_password(request):
    step = "username"

    if request.method == "POST":
        action = request.POST.get("action")

        # -----------------------------------
        # STEP 1: SEND OTP
        # -----------------------------------
        if action == "send_otp":
            username = request.POST.get("username")

            try:
                user = User.objects.get(
                    username=username,
                    is_superuser=True
                )

            except User.DoesNotExist:
                return render(
                    request,
                    "admin_forgot_password.html",
                    {
                        "step": "username",
                        "error": "Admin account not found."
                    }
                )

            # Generate 6-digit OTP
            otp = str(random.randint(100000, 999999))

            # Save OTP temporarily in session
            request.session["reset_otp"] = otp
            request.session["reset_username"] = username

            request.session["otp_expiry"] = (
                timezone.now() +
                timedelta(minutes=10)
            ).timestamp()

            # Print OTP in Django terminal
            print("\n")
            print("=" * 40)
            print("ADMIN PASSWORD RESET OTP")
            print("=" * 40)
            print(f"Username: {username}")
            print(f"OTP: {otp}")
            print("=" * 40)
            print("\n")

            return render(
                request,
                "accounts/admin_forgot_password.html",
                {
                    "step": "otp",
                    "username": username
                }
            )

        # -----------------------------------
        # STEP 2: VERIFY OTP
        # -----------------------------------
        elif action == "verify_otp":
            username = request.POST.get("username")
            entered_otp = request.POST.get("otp")

            saved_otp = request.session.get("reset_otp")
            saved_username = request.session.get("reset_username")
            otp_expiry = request.session.get("otp_expiry")

            # Validate session
            if (
                not saved_otp
                or not saved_username
                or not otp_expiry
            ):
                return redirect("admin_forgot_password")

            # Check OTP expiry
            if timezone.now().timestamp() > otp_expiry:

                request.session.pop("reset_otp", None)
                request.session.pop("reset_username", None)
                request.session.pop("otp_expiry", None)

                return render(
                    request,
                    "accounts/admin_forgot_password.html",
                    {
                        "step": "username",
                        "error": "OTP expired. Please request a new one."
                    }
                )

            # Check username and OTP
            if (
                username != saved_username
                or entered_otp != saved_otp
            ):
                return render(
                    request,
                    "accounts/admin_forgot_password.html",
                    {
                        "step": "otp",
                        "username": saved_username,
                        "error": "Invalid OTP."
                    }
                )

            # Mark OTP as verified
            request.session["otp_verified"] = True

            return render(
                request,
                "accounts/admin_forgot_password.html",
                {
                    "step": "password",
                    "username": saved_username
                }
            )

        # -----------------------------------
        # STEP 3: RESET PASSWORD
        # -----------------------------------
        elif action == "reset_password":

            username = request.POST.get("username")
            new_password = request.POST.get("new_password")
            confirm_password = request.POST.get(
                "confirm_password"
            )

            # Security check
            if (
                not request.session.get("otp_verified")
                or username != request.session.get(
                    "reset_username"
                )
            ):
                return redirect("admin_forgot_password")

            # Check passwords match
            if new_password != confirm_password:
                return render(
                    request,
                    "accounts/admin_forgot_password.html",
                    {
                        "step": "password",
                        "username": username,
                        "error": "Passwords do not match."
                    }
                )

            try:
                user = User.objects.get(
                    username=username,
                    is_superuser=True
                )

            except User.DoesNotExist:
                return redirect("admin_forgot_password")

            # Change password securely
            user.set_password(new_password)
            user.save()

            # Remove all reset data from session
            request.session.pop("reset_otp", None)
            request.session.pop("reset_username", None)
            request.session.pop("otp_expiry", None)
            request.session.pop("otp_verified", None)

            # Redirect back to login
            return redirect("admin_login")

    # -----------------------------------
    # DEFAULT PAGE
    # -----------------------------------
    return render(
        request,
        "accounts/admin_forgot_password.html",
        {
            "step": step
        }
    )
# ---------------------------------
# SALES MEMBER LOGIN
# ---------------------------------
@never_cache
def sales_login(request):

    if request.method == "POST":

        username = request.POST.get("username")
        password = request.POST.get("password")

        user = authenticate(
            request,
            username=username,
            password=password
        )

        if user is not None and not user.is_superuser:

            login(request, user)

            return redirect("sales_dashboard")

        return render(
            request,
            "accounts/sales_login.html",
            {
                "error": "Invalid Username or Password"
            }
        )

    return render(
        request,
        "accounts/sales_login.html"
    )


# ---------------------------------
# LOGOUT
# ---------------------------------
@login_required
def logout_view(request):

    logout(request)

    return redirect("/")

@login_required
@user_passes_test(is_admin)
def sales_member_list(request):

    members = User.objects.filter(
        is_superuser=False,
        is_active=True
    )

    for member in members:

        member.pending_count = Order.objects.filter(
            sales_member=member,
            status="Pending"
        ).count()

        member.dispatched_count = Order.objects.filter(
            sales_member=member,
            status="Completed"
        ).count()

        member.out_of_stock_count = Order.objects.filter(
            sales_member=member,
            status="Out of Stock"
        ).count()

        member.total_sales = Order.objects.filter(
            sales_member=member,
            status="Completed"
        ).aggregate(
            total=Sum("amount")
        )["total"] or 0

    return render(
    request,
    "accounts/sales_members.html",
    {
        "members": members
    }
)

# ---------------------------------
# ADD SALES MEMBER
# ---------------------------------
@login_required
@user_passes_test(is_admin)
def sales_member_create(request):

    if request.method == "POST":

        form = SalesMemberForm(request.POST)

        if form.is_valid():

            user = form.save(commit=False)

            user.set_password(
                form.cleaned_data["password"]
            )

            user.is_superuser = False
            user.is_staff = False
            user.is_active = True

            user.save()

            return redirect("sales_list")

    else:

        form = SalesMemberForm()

    return render(
        request,
        "accounts/add_sales.html",
        {
            "form": form
        }
    )


# ---------------------------------
# DELETE SALES MEMBER
# ---------------------------------
@login_required
@user_passes_test(is_admin)
def delete_sales_member(request, id):

    user = get_object_or_404(
        User,
        id=id,
        is_superuser=False
    )

    # Soft delete
    user.is_active = False
    user.save()

    return redirect("sales_list")

@login_required
@user_passes_test(is_admin)
def sales_member_sales_detail(request, member_id):

    member = get_object_or_404(
        User,
        id=member_id,
        is_superuser=False
    )

    all_orders = Order.objects.filter(
    sales_member=member
).order_by("-updated_at")

    pending_count = all_orders.filter(
        status="Pending"
    ).count()

    dispatched_count = all_orders.filter(
        status="Completed"
    ).count()

    out_of_stock_count = all_orders.filter(
        status="Out of Stock"
    ).count()

    total_sales = all_orders.filter(
        status="Completed"
    ).aggregate(
        total=Sum("amount")
    )["total"] or 0

    return render(
        request,
        "accounts/sales_member_sales_detail.html",
        {
            "member": member,
            "orders": all_orders,
            "pending_count": pending_count,
            "dispatched_count": dispatched_count,
            "out_of_stock_count": out_of_stock_count,
            "total_sales": total_sales,
        }
    )
@login_required
@user_passes_test(is_admin)
def export_sales_member_excel(request, member_id):

    member = get_object_or_404(
        User,
        id=member_id,
        is_superuser=False
    )

    # Only completed / dispatched orders
    orders = Order.objects.filter(
        sales_member=member,
        status="Completed"
    ).order_by("id")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sales Report"

    # ==========================================
    # TITLE
    # ==========================================

    worksheet.merge_cells("A1:C1")

    title_cell = worksheet["A1"]
    title_cell.value = "SALES REPORT"

    title_cell.font = Font(
        size=20,
        bold=True,
        color="FFFFFF"
    )

    title_cell.fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    worksheet.row_dimensions[1].height = 35


    # ==========================================
    # SALES MEMBER NAME
    # ==========================================

    worksheet.merge_cells("A2:C2")

    name_cell = worksheet["A2"]

    name_cell.value = (
        f"Sales Member: "
        f"{member.first_name} {member.last_name}"
    )

    name_cell.font = Font(
        size=14,
        bold=True
    )

    name_cell.fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7"
    )

    name_cell.alignment = Alignment(
        horizontal="left",
        vertical="center"
    )

    worksheet.row_dimensions[2].height = 25


    # ==========================================
    # TABLE HEADER
    # ==========================================

    headers = [
        "Sl No",
        "Company Name",
        "Sales Amount"
    ]

    for column, header in enumerate(headers, 1):

        cell = worksheet.cell(
            row=4,
            column=column,
            value=header
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    worksheet.row_dimensions[4].height = 25


    # ==========================================
    # ORDER DATA
    # ==========================================

    total_sales = 0

    for index, order in enumerate(orders, 1):

        row = index + 4

        worksheet.cell(
            row=row,
            column=1,
            value=index
        )

        worksheet.cell(
            row=row,
            column=2,
            value=order.company_name
        )

        amount = order.amount or 0

        worksheet.cell(
            row=row,
            column=3,
            value=float(amount)
        )

        total_sales += amount

        # Row alignment
        for column in range(1, 4):

            cell = worksheet.cell(
                row=row,
                column=column
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

            cell.border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin")
            )


    # ==========================================
    # TOTAL SALES
    # ==========================================

    total_row = len(orders) + 6

    worksheet.merge_cells(
        start_row=total_row,
        start_column=1,
        end_row=total_row,
        end_column=2
    )

    total_label = worksheet.cell(
        row=total_row,
        column=1
    )

    total_label.value = "TOTAL SALES"

    total_label.font = Font(
        bold=True,
        size=14
    )

    total_label.alignment = Alignment(
        horizontal="right"
    )

    total_amount = worksheet.cell(
        row=total_row,
        column=3,
        value=float(total_sales)
    )

    total_amount.font = Font(
        bold=True,
        size=14
    )

    total_amount.alignment = Alignment(
        horizontal="center"
    )

    for column in range(1, 4):

        cell = worksheet.cell(
            row=total_row,
            column=column
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="E2F0D9"
        )

        cell.border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )


    # ==========================================
    # COLUMN WIDTH
    # ==========================================

    worksheet.column_dimensions["A"].width = 12
    worksheet.column_dimensions["B"].width = 35
    worksheet.column_dimensions["C"].width = 20


    # ==========================================
    # CURRENCY FORMAT
    # ==========================================

    for row in range(5, total_row + 1):

        worksheet.cell(
            row=row,
            column=3
        ).number_format = '₹#,##0.00'


    # ==========================================
    # DOWNLOAD EXCEL
    # ==========================================

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )
    )

    response[
        "Content-Disposition"
    ] = (
        f'attachment; filename="'
        f'{member.username}_sales_report.xlsx"'
    )

    workbook.save(response)

    return response